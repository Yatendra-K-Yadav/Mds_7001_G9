"""
Build the analysis table.

Reads the five raw sources, reshapes each into a tidy frame keyed on
(council, year), joins them, and derives the project's core metrics.

    python src/ingest.py

Outputs
-------
data/interim/    one tidy CSV per source, plus the council-name crosswalk
data/processed/  analysis_table.csv  -- one row per council per year
"""

import io
import re
import unicodedata
import urllib.request

import numpy as np
import openpyxl
import pandas as pd

from src.config import (ADULT_AGE_FROM, BASE_YEAR, CPI_API_URL, GAMBLING_CSV,
                    INCOME_XLSX, INTERIM, POPULATION_XLSX, PROCESSED,
                    QLD_LGA_CODES, SEIFA_XLSX)


# --------------------------------------------------------------------------
# Council names are spelled differently by each publisher: the gambling data
# uses "MCKINLAY", the ABS uses "McKinlay", and SEIFA appends "(Qld)" where a
# name is ambiguous across states. Normalising to a common key lets the five
# sources join on one column.
# --------------------------------------------------------------------------
def normalise_lga(name):
    """Reduce a council name to a join key: uppercase, no suffix, no punctuation."""
    text = str(name).strip()
    text = re.sub(r"\s*\((?:Qld|S|C|R|T|DC|M|A)\)\s*$", "", text, flags=re.I)
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    text = re.sub(r"[^A-Za-z0-9 ]", " ", text)
    return re.sub(r"\s+", " ", text).strip().upper()


def load_gambling():
    """Monthly gaming machine data, one row per council per month."""
    df = pd.read_csv(GAMBLING_CSV, encoding="utf-8-sig").rename(columns={
        "LGA Region": "lga_raw",
        "Approved Sites": "approved_sites",
        "Operational Sites": "operational_sites",
        "Approved EGMs": "approved_egms",
        "Operational EGMs": "operational_egms",
        "Metered Win": "metered_win",
    })
    df["date"] = pd.to_datetime(df["Month Year"], format="%B %Y")
    df["year"] = df["date"].dt.year
    df["lga_key"] = df["lga_raw"].map(normalise_lga)

    # A blank metered_win means the figure was withheld: the regulator
    # suppresses any council-month with fewer than six operating venues.
    df["suppressed"] = df["metered_win"].isna()

    print(f"[gambling]   {len(df):,} monthly rows, {df.year.min()}-{df.year.max()}, "
          f"{df.lga_key.nunique()} councils")
    print(f"[gambling]   {df.suppressed.sum():,} rows suppressed "
          f"({100 * df.suppressed.mean():.1f}%)")
    return df


def aggregate_gambling_to_years(monthly):
    """Collapse months into council-years, flagging any that are incomplete."""
    annual = (monthly
              .groupby(["lga_key", "lga_raw", "year"])
              .agg(months=("date", "nunique"),
                   months_suppressed=("suppressed", "sum"),
                   metered_win=("metered_win", "sum"),
                   operational_egms=("operational_egms", "mean"),
                   operational_sites=("operational_sites", "mean"))
              .reset_index())

    # Only a full year with no withheld months can be compared against another.
    annual["complete"] = (annual["months"] == 12) & (annual["months_suppressed"] == 0)

    print(f"[gambling]   {len(annual):,} council-years, "
          f"{annual.complete.sum():,} complete ({100 * annual.complete.mean():.1f}%)")
    return annual


def load_population():
    """Queensland adult population by council and year, from ABS age bands."""
    workbook = openpyxl.load_workbook(POPULATION_XLSX, read_only=True, data_only=True)
    rows = list(workbook["Table 3"].iter_rows(min_row=5, values_only=True))
    workbook.close()

    band_labels = [str(cell).strip() for cell in rows[0][5:] if cell]
    header = rows[1]
    qld_rows = [r for r in rows[2:] if r[0] and str(r[2]).strip() == "Queensland"]

    columns = [str(c).strip() if c else f"col{i}" for i, c in enumerate(header)]
    df = pd.DataFrame(qld_rows, columns=columns)
    df.columns = columns[:5] + band_labels + columns[5 + len(band_labels):]
    df = df.rename(columns={columns[0]: "year",
                            columns[3]: "abs_lga_code",
                            columns[4]: "abs_lga_name"})

    # "Total persons" sits alongside the age bands and must not be summed with them.
    age_bands = [b for b in band_labels if b != "Total persons"]
    adult_bands = [b for b in age_bands
                   if re.match(r"^(\d+)", b)
                   and int(re.match(r"^(\d+)", b).group(1)) >= ADULT_AGE_FROM]

    df["adults"] = df[adult_bands].apply(pd.to_numeric, errors="coerce").sum(axis=1)
    df["pop_total"] = pd.to_numeric(df["Total persons"], errors="coerce")

    summed = df[age_bands].apply(pd.to_numeric, errors="coerce").sum(axis=1)
    assert (abs(summed - df["pop_total"]) < 2).all(), \
        "age bands do not sum to the published total -- check the band labels"

    df["year"] = df["year"].astype(int)
    df["lga_key"] = df["abs_lga_name"].map(normalise_lga)

    print(f"[population] {len(df):,} QLD council-years, {df.year.min()}-{df.year.max()}, "
          f"adults counted from age {ADULT_AGE_FROM}")
    return df[["year", "abs_lga_code", "abs_lga_name", "lga_key", "adults", "pop_total"]]


def load_cpi():
    """Brisbane All-groups CPI, quarterly from the ABS API, averaged to years."""
    request = urllib.request.Request(
        CPI_API_URL, headers={"Accept": "application/vnd.sdmx.data+csv"})
    quarterly = pd.read_csv(io.BytesIO(urllib.request.urlopen(request, timeout=90).read()))
    quarterly["year"] = quarterly["TIME_PERIOD"].str[:4].astype(int)

    cpi = quarterly.groupby("year")["OBS_VALUE"].mean().rename("cpi_index").reset_index()
    base_index = float(cpi.loc[cpi.year == BASE_YEAR, "cpi_index"].iloc[0])

    # Multiply a nominal dollar amount by this to express it in BASE_YEAR money.
    cpi["cpi_factor"] = base_index / cpi["cpi_index"]

    print(f"[cpi]        Brisbane {cpi.year.min()}-{cpi.year.max()}, "
          f"base {BASE_YEAR} (index {base_index:.1f})")
    return cpi


def load_seifa():
    """SEIFA 2021 Index of Relative Socio-economic Disadvantage, QLD councils."""
    workbook = openpyxl.load_workbook(SEIFA_XLSX, read_only=True, data_only=True)
    records = []
    for row in workbook["Table 1"].iter_rows(min_row=7, max_col=4, values_only=True):
        try:
            code = int(row[0])
        except (TypeError, ValueError):
            continue
        if QLD_LGA_CODES[0] <= code < QLD_LGA_CODES[1] and row[2] is not None:
            records.append({"abs_lga_code": code,
                            "seifa_lga_name": str(row[1]).strip(),
                            "irsd_score": float(row[2]),
                            "irsd_decile": int(row[3])})
    workbook.close()

    seifa = pd.DataFrame(records)
    seifa["lga_key"] = seifa["seifa_lga_name"].map(normalise_lga)
    print(f"[seifa]      {len(seifa)} QLD councils, IRSD "
          f"{seifa.irsd_score.min():.0f}-{seifa.irsd_score.max():.0f}")
    return seifa


def load_income():
    """Median total income by QLD council and financial year (ABS, from ATO data)."""
    workbook = openpyxl.load_workbook(INCOME_XLSX, read_only=True, data_only=True)
    rows = list(workbook["Table 1.5"].iter_rows(min_row=6, values_only=True))
    workbook.close()

    # Row 1 names each measure block, row 2 the years within it; forward-fill
    # the block name so every column becomes "Measure|Year".
    blocks = [str(c).strip() if c else "" for c in rows[0]]
    years = [str(c).strip() if c else "" for c in rows[1]]
    columns, current_block = [], ""
    for block, year in zip(blocks, years):
        if block:
            current_block = block
        columns.append(f"{current_block}|{year}" if year not in ("", "None")
                       else (block or year))

    df = pd.DataFrame(rows[2:], columns=columns).rename(
        columns={columns[0]: "abs_lga_code", columns[1]: "abs_lga_name"})
    median_columns = [c for c in columns if c.startswith("Median ($)|")]

    qld = df[pd.to_numeric(df["abs_lga_code"], errors="coerce")
             .between(*[c - 1 for c in QLD_LGA_CODES])].copy()
    qld["lga_key"] = qld["abs_lga_name"].map(normalise_lga)

    income = qld.melt(id_vars=["lga_key"], value_vars=median_columns,
                      var_name="column", value_name="median_income")
    income["financial_year"] = income["column"].str.split("|").str[1]
    income["year"] = income["financial_year"].str[:4].astype(int) + 1
    income["median_income"] = pd.to_numeric(income["median_income"], errors="coerce")
    income = income[["lga_key", "year", "median_income"]]

    print(f"[income]     {qld.lga_key.nunique()} QLD councils, "
          f"years {sorted(income.year.unique())}")
    return income


def derive_metrics(df):
    """Add the project's core per-adult and per-machine measures."""
    df["real_win"] = df["metered_win"] * df["cpi_factor"]
    df["loss_per_adult"] = df["real_win"] / df["adults"]
    df["loss_per_egm"] = df["real_win"] / df["operational_egms"]
    df["egms_per_1000_adults"] = 1000 * df["operational_egms"] / df["adults"]
    df["irsd_quintile"] = pd.qcut(
        df["irsd_score"], 5,
        labels=["Q1 most disadv", "Q2", "Q3", "Q4", "Q5 least disadv"])
    return df


def main():
    print("Loading sources")
    monthly = load_gambling()
    annual = aggregate_gambling_to_years(monthly)
    population = load_population()
    cpi = load_cpi()
    seifa = load_seifa()
    income = load_income()

    print("\nJoining on (council, year)")
    analysis = (annual
                .merge(population, on=["lga_key", "year"], how="left")
                .merge(cpi, on="year", how="left")
                .merge(seifa[["lga_key", "abs_lga_code", "irsd_score", "irsd_decile"]],
                       on="lga_key", how="left", suffixes=("", "_seifa"))
                .merge(income, on=["lga_key", "year"], how="left"))
    analysis = derive_metrics(analysis)

    INTERIM.mkdir(parents=True, exist_ok=True)
    PROCESSED.mkdir(parents=True, exist_ok=True)

    crosswalk = (annual[["lga_key", "lga_raw"]].drop_duplicates()
                 .merge(seifa[["lga_key", "seifa_lga_name", "abs_lga_code"]],
                        on="lga_key", how="left")
                 .sort_values("lga_raw"))

    for frame, path in [(monthly, INTERIM / "gambling_monthly.csv"),
                        (population, INTERIM / "population_qld.csv"),
                        (cpi, INTERIM / "cpi_brisbane_annual.csv"),
                        (seifa, INTERIM / "seifa_qld.csv"),
                        (income, INTERIM / "income_qld.csv"),
                        (crosswalk, INTERIM / "lga_crosswalk.csv"),
                        (analysis, PROCESSED / "analysis_table.csv")]:
        frame.to_csv(path, index=False)

    unmatched_population = analysis.adults.isna().sum()
    print(f"\n[join]       analysis table {analysis.shape[0]:,} rows x {analysis.shape[1]} cols")
    print(f"[join]       {unmatched_population} rows without population "
          f"(population series ends {population.year.max()})")
    print(f"[join]       {analysis.irsd_score.isna().sum()} rows without SEIFA")
    print(f"\nWrote {PROCESSED / 'analysis_table.csv'}")


if __name__ == "__main__":
    main()
